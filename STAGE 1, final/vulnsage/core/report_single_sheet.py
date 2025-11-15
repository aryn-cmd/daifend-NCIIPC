"""
core/report_single_sheet.py
--------------------------------
Create a single-sheet PS01-style report (XLSX or CSV fallback) from a merged-results JSON.

Usage:
    python report_single_sheet.py --input <merged_results.json> --name <AIGR-name> --out-suffix single

This script purposely keeps behavior simple and independent from the main report_builder
so you can generate single-sheet files for specific stacks (e.g., stack3 and stack4).
"""
import os
import json
import csv
import argparse
from typing import List, Dict, Any

HEADERS = [
    "SNo",
    "Primary Language of Benchmark",
    "Vulnerability",
    "CVE ID",
    "Severity",
    "Common Weakness Enumeration (CWE) Id",
    "file name with path",
    "line number",
    "Code Snippet at the line",
]


def load_findings(path: str) -> List[Dict[str, Any]]:
    with open(path, 'r', encoding='utf8') as fh:
        data = json.load(fh)
    if isinstance(data, dict):
        for key in ("findings", "results", "detections", "flat"):
            if key in data and isinstance(data[key], list):
                return data[key]
        for v in data.values():
            if isinstance(v, list):
                return v
        raise ValueError("Findings JSON dict does not contain a list of findings")
    if not isinstance(data, list):
        raise ValueError("Findings JSON must be a list or a dict containing a list")
    return data


def _extract_cve_from_det(det: Dict[str, Any]) -> str:
    if det.get('cve'):
        return str(det.get('cve'))
    cm = det.get('cve_matches') or det.get('cves')
    if isinstance(cm, list) and cm:
        first = cm[0]
        return str(first.get('cve') or first.get('id') or first)
    return ""


def _get_severity(det: Dict[str, Any]) -> str:
    if det is None:
        return "Unknown"
    sev = det.get('severity')
    if sev:
        return str(sev)
    for key in ('accuracy', 'confidence', 'score', 'probability'):
        v = det.get(key)
        if v is None:
            continue
        try:
            fv = float(v)
        except Exception:
            continue
        if fv >= 0.85:
            return 'High'
        if fv >= 0.6:
            return 'Medium'
        if fv >= 0.3:
            return 'Low'
        return 'Info'
    return 'Unknown'


def _format_vuln_field(det: Dict[str, Any]) -> str:
    title = det.get('vulnerability') or det.get('rule_id') or det.get('id') or 'Vulnerability'
    desc = det.get('description') or det.get('message') or det.get('fix') or ''
    if desc:
        return f"{title} — {desc}"
    return title


def _extract_code_snippet(det: Dict[str, Any]) -> str:
    path = det.get('file')
    line_no = det.get('line')
    if not path or not line_no:
        return det.get('evidence', '')
    try:
        with open(path, 'r', encoding='utf8', errors='ignore') as fh:
            lines = fh.readlines()
        idx = int(line_no) - 1
        if 0 <= idx < len(lines):
            return lines[idx].strip()
        return det.get('evidence', '')
    except Exception:
        return det.get('evidence', '')


def _extract_cwe_from_det(det: Dict[str, Any]) -> str:
    if det.get('cwe'):
        return str(det.get('cwe'))
    if det.get('cwe_id'):
        return str(det.get('cwe_id'))
    return ""
    try:
        with open(path, 'r', encoding='utf8', errors='ignore') as fh:
            lines = fh.readlines()
        idx = int(line_no) - 1
        if 0 <= idx < len(lines):
            return lines[idx].strip()
        return det.get('evidence', '')
    except Exception:
        return det.get('evidence', '')


def filter_vulns(findings: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for f in findings:
        has_id = bool(f.get('vulnerability') or f.get('rule_id') or f.get('id'))
        numeric_score = None
        for key in ('accuracy', 'confidence', 'score', 'probability'):
            if key in f:
                try:
                    numeric_score = float(f[key])
                    break
                except Exception:
                    pass
        if has_id or (numeric_score is not None and numeric_score >= 0.6):
            out.append(f)
    return out


def write_single_sheet(findings: List[Dict[str, Any]], out_name: str):
    os.makedirs('output', exist_ok=True)
    out_xlsx = os.path.join('output', out_name)

    # Try openpyxl
    try:
        import importlib
        openpyxl = importlib.import_module('openpyxl')
        Workbook = getattr(openpyxl, 'Workbook')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Findings'
        ws.append(HEADERS)
        rowno = 1
        for det in findings:
            # infer language from file extension if missing or 'unknown'
            lang = (det.get('language') or det.get('lang') or '').lower()
            if not lang or lang == 'unknown':
                fp = det.get('file', '')
                ext = os.path.splitext(fp)[1].lower()
                ext_map = {
                    '.py': 'python', '.java': 'java', '.js': 'javascript', '.ts': 'typescript',
                    '.c': 'c', '.cpp': 'cpp', '.cc': 'cpp', '.h': 'c', '.cs': 'csharp', '.php': 'php'
                }
                lang = ext_map.get(ext, lang or 'unknown')
            rowno += 1
            cve = _extract_cve_from_det(det)
            cwe_field = det.get('cwe') or det.get('cwe_id') or _extract_cwe_from_det(det) or ""
            cve_field = cve or ""
            ws.append([
                rowno - 1,
                lang.title(),
                _format_vuln_field(det),
                cve_field,
                _get_severity(det),
                cwe_field,
                det.get('file', ''),
                det.get('line', ''),
                _extract_code_snippet(det),
            ])
        # modest column widths
        try:
            from openpyxl.utils import get_column_letter
            for i in range(1, len(HEADERS) + 1):
                ws.column_dimensions[get_column_letter(i)].width = 22
        except Exception:
            pass
        wb.save(out_xlsx)
        return out_xlsx
    except ModuleNotFoundError:
        # CSV fallback
        out_csv = out_xlsx.replace('.xlsx', '.csv')
        with open(out_csv, 'w', encoding='utf8', newline='') as fh:
            writer = csv.writer(fh)
            writer.writerow(HEADERS)
            for i, det in enumerate(findings, start=1):
                cve = _extract_cve_from_det(det)
                cwe_field = det.get('cwe') or det.get('cwe_id') or _extract_cwe_from_det(det) or ""
                cve_field = cve or ""
                lang = (det.get('language') or det.get('lang') or '').lower()
                if not lang or lang == 'unknown':
                    fp = det.get('file', '')
                    ext = os.path.splitext(fp)[1].lower()
                    ext_map = {'.py': 'python', '.java': 'java', '.js': 'javascript', '.ts': 'typescript', '.c': 'c', '.cpp': 'cpp', '.cc': 'cpp', '.h': 'c', '.cs': 'csharp', '.php': 'php'}
                    lang = ext_map.get(ext, lang or 'unknown')
                writer.writerow([
                    i,
                    lang.title(),
                    _format_vuln_field(det),
                    cve_field,
                    _get_severity(det),
                    cwe_field,
                    det.get('file', ''),
                    det.get('line', ''),
                    _extract_code_snippet(det),
                ])
        return out_csv


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True, help='Merged results JSON path')
    parser.add_argument('--name', required=True, help='Output base name (e.g., AIGR-S66944_stack3)')
    parser.add_argument('--out-suffix', default='single', help='Suffix to append to filename')
    args = parser.parse_args()

    findings = load_findings(args.input)
    findings = filter_vulns(findings)
    if not findings:
        print('[WARN] No vulnerabilities after filtering; no report created')
        return
    out_name = f"GC_PS_01_{args.name}_{args.out_suffix}.xlsx"
    out_path = write_single_sheet(findings, out_name)
    print(f'[INFO] Single-sheet report created: {out_path}')


if __name__ == '__main__':
    main()
