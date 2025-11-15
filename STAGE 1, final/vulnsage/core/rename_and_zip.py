import os
import zipfile

out_dir = r"d:\NCIIPC\STAGE 1, final\output"
aigr_id = 'S66944'
old_prefix = f'AIGR-{aigr_id}_'
new_prefix = 'GC_PS_01_DAIFEND_'

renamed = []
for fn in os.listdir(out_dir):
    if fn.startswith(old_prefix) and fn.lower().endswith('.xlsx'):
        old = os.path.join(out_dir, fn)
        rest = fn.split('_', 1)[1] if '_' in fn else fn
        new_name = new_prefix + rest
        new = os.path.join(out_dir, new_name)
        try:
            os.replace(old, new)
            renamed.append(new)
            print(f'Renamed: {old} -> {new}')
        except Exception as e:
            print('Rename failed:', old, e)

# Recreate zip with GC_PS_01 files
zip_path = os.path.join(out_dir, f'AIGR-{aigr_id}.zip')
with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
    for f in os.listdir(out_dir):
        if f.startswith('GC_PS_01_DAIFEND_') and f.lower().endswith('.xlsx'):
            path = os.path.join(out_dir, f)
            zf.write(path, arcname=f)
            print('Zipped:', f)

print('\nFinal directory listing:')
for f in sorted(os.listdir(out_dir)):
    print(f)

# Quick verify one workbook
try:
    from openpyxl import load_workbook
    sample = os.path.join(out_dir, 'GC_PS_01_DAIFEND_S1_4.xlsx')
    if os.path.exists(sample):
        wb = load_workbook(sample, read_only=True)
        print('\nSample workbook sheets:', wb.sheetnames)
        ws = wb[wb.sheetnames[0]]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        print('Sample header:', header)
        row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
        print('Sample row:', row)
    else:
        print('Sample workbook not found:', sample)
except Exception as e:
    print('openpyxl not available or failed to read workbook:', e)
