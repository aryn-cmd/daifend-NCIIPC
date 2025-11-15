import sys
from pathlib import Path

try:
    import openpyxl
except Exception as e:
    print("openpyxl not installed:", e)
    sys.exit(2)

if len(sys.argv) < 2:
    print("Usage: inspect_xlsx.py <path-to-xlsx>")
    sys.exit(1)

p = Path(sys.argv[1])
if not p.exists():
    print(f"File not found: {p}")
    sys.exit(1)

wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print(f"Workbook: {p}\nSheets: {wb.sheetnames}\n")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"--- Sheet: {name} (max_row={ws.max_row}) ---")
    rows = ws.iter_rows(values_only=True)
    # print header
    try:
        header = next(rows)
    except StopIteration:
        print("(empty sheet)")
        continue
    print("Header:", header)
    count = 0
    for r in rows:
        print(r)
        count += 1
        if count >= 5:
            break
    print()

print("Done.")
